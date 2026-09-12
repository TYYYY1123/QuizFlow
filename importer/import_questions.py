import os
import csv
import random

from openpyxl import load_workbook
from docx import Document
from PyPDF2 import PdfReader


# ==================================================
# 路径
# ==================================================

PROJECT_ROOT = os.path.dirname(
    os.path.dirname(os.path.abspath(__file__))
)

INPUT_DIR = os.path.join(PROJECT_ROOT, "input")
OUTPUT_DIR = os.path.join(PROJECT_ROOT, "data")
OUTPUT_FILE = os.path.join(OUTPUT_DIR, "questions.txt")


# ==================================================
# 第一版：普通复习资料自动出题规则
# 后面会升级成真正的智能出题模块
# ==================================================

RULES = [
    {
        "keyword": "printf",
        "options": ["printf", "scanf", "main", "char"]
    },
    {
        "keyword": "scanf",
        "options": ["scanf", "printf", "main", "strlen"]
    },
    {
        "keyword": "char",
        "options": ["char", "int", "float", "double"]
    },
    {
        "keyword": "main",
        "options": ["main", "printf", "scanf", "start"]
    },
    {
        "keyword": "0",
        "options": ["0", "1", "-1", "2"]
    }
]


# ==================================================
# 普通文本 → 自动生成选择题
# ==================================================

def make_question_from_sentence(sentence):

    sentence = sentence.strip()

    if not sentence:
        return None

    for rule in RULES:

        keyword = rule["keyword"]

        if keyword in sentence:

            question_text = sentence.replace(
                keyword,
                "____",
                1
            )

            options = rule["options"][:]

            random.shuffle(options)

            correct_index = options.index(keyword)

            answer = chr(
                ord("A") + correct_index
            )

            return {
                "question": question_text,
                "options": options,
                "answer": answer
            }

    return None


def sentences_to_questions(sentences):

    questions = []

    for sentence in sentences:

        question = make_question_from_sentence(
            sentence
        )

        if question is not None:
            questions.append(question)

    return questions


# ==================================================
# TXT
# ==================================================

def read_txt_file(file_path):

    sentences = []

    with open(
        file_path,
        "r",
        encoding="utf-8"
    ) as f:

        for line in f:

            line = line.strip()

            if line:
                sentences.append(line)

    return sentences_to_questions(sentences)


# ==================================================
# Excel XLSX
# 标准格式：
# question | A | B | C | D | answer
# ==================================================

def read_xlsx_file(file_path):

    questions = []

    workbook = load_workbook(
        file_path,
        data_only=True
    )

    sheet = workbook.active

    rows = list(
        sheet.iter_rows(values_only=True)
    )

    if not rows:
        return questions

    headers = []

    for value in rows[0]:

        if value is None:
            headers.append("")
        else:
            headers.append(
                str(value).strip().lower()
            )

    required = [
        "question",
        "a",
        "b",
        "c",
        "d",
        "answer"
    ]

    for name in required:

        if name not in headers:

            print(
                "Excel格式错误，缺少列：{}".format(name)
            )

            return questions

    qi = headers.index("question")
    ai = headers.index("a")
    bi = headers.index("b")
    ci = headers.index("c")
    di = headers.index("d")
    ans_i = headers.index("answer")

    for row in rows[1:]:

        if not row:
            continue

        if row[qi] is None:
            continue

        values = [
            row[ai],
            row[bi],
            row[ci],
            row[di],
            row[ans_i]
        ]

        if None in values:
            continue

        questions.append({
            "question":
                str(row[qi]).strip(),

            "options": [
                str(row[ai]).strip(),
                str(row[bi]).strip(),
                str(row[ci]).strip(),
                str(row[di]).strip()
            ],

            "answer":
                str(row[ans_i])
                .strip()
                .upper()
        })

    return questions


# ==================================================
# CSV
# 格式和 Excel 一样
# ==================================================

def read_csv_file(file_path):

    questions = []

    with open(
        file_path,
        "r",
        encoding="utf-8-sig",
        newline=""
    ) as f:

        reader = csv.DictReader(f)

        if reader.fieldnames is None:
            return questions

        field_map = {}

        for name in reader.fieldnames:

            if name is not None:

                field_map[
                    name.strip().lower()
                ] = name

        required = [
            "question",
            "a",
            "b",
            "c",
            "d",
            "answer"
        ]

        for name in required:

            if name not in field_map:

                print(
                    "CSV格式错误，缺少列：{}".format(name)
                )

                return questions

        for row in reader:

            question_text = row[
                field_map["question"]
            ]

            if not question_text:
                continue

            questions.append({
                "question":
                    question_text.strip(),

                "options": [
                    row[field_map["a"]].strip(),
                    row[field_map["b"]].strip(),
                    row[field_map["c"]].strip(),
                    row[field_map["d"]].strip()
                ],

                "answer":
                    row[field_map["answer"]]
                    .strip()
                    .upper()
            })

    return questions


# ==================================================
# Word DOCX
# 第一版：读取普通段落
# ==================================================

def read_docx_file(file_path):

    document = Document(file_path)

    sentences = []

    for paragraph in document.paragraphs:

        text = paragraph.text.strip()

        if text:
            sentences.append(text)

    return sentences_to_questions(sentences)


# ==================================================
# PDF
# 第一版：读取文字型 PDF
# 扫描图片 PDF 暂时无法识别
# ==================================================

def read_pdf_file(file_path):

    sentences = []

    reader = PdfReader(file_path)

    for page in reader.pages:

        text = page.extract_text()

        if not text:
            continue

        for line in text.splitlines():

            line = line.strip()

            if line:
                sentences.append(line)

    return sentences_to_questions(sentences)


# ==================================================
# 扫描 input 文件夹
# ==================================================

def load_all_input_files():

    all_questions = []

    if not os.path.exists(INPUT_DIR):
        os.makedirs(INPUT_DIR)

    for filename in os.listdir(INPUT_DIR):

        # 忽略 Excel / Word 临时文件
        if filename.startswith("~$"):
            continue

        file_path = os.path.join(
            INPUT_DIR,
            filename
        )

        lower_name = filename.lower()

        try:

            if lower_name.endswith(".txt"):

                print(
                    "读取 TXT：{}".format(filename)
                )

                all_questions.extend(
                    read_txt_file(file_path)
                )

            elif lower_name.endswith(".xlsx"):

                print(
                    "读取 Excel：{}".format(filename)
                )

                all_questions.extend(
                    read_xlsx_file(file_path)
                )

            elif lower_name.endswith(".csv"):

                print(
                    "读取 CSV：{}".format(filename)
                )

                all_questions.extend(
                    read_csv_file(file_path)
                )

            elif lower_name.endswith(".docx"):

                print(
                    "读取 Word：{}".format(filename)
                )

                all_questions.extend(
                    read_docx_file(file_path)
                )

            elif lower_name.endswith(".pdf"):

                print(
                    "读取 PDF：{}".format(filename)
                )

                all_questions.extend(
                    read_pdf_file(file_path)
                )

            else:

                print(
                    "暂不支持：{}".format(filename)
                )

        except Exception as e:

            print(
                "读取失败：{}".format(filename)
            )

            print(
                "原因：{}".format(e)
            )

    return remove_duplicates(
        all_questions
    )


# ==================================================
# 去重
# ==================================================

def remove_duplicates(questions):

    result = []

    seen = set()

    for q in questions:

        key = (
            q["question"],
            tuple(q["options"]),
            q["answer"]
        )

        if key not in seen:

            seen.add(key)

            result.append(q)

    return result


# ==================================================
# 保存统一题库
# ==================================================

def save_questions(questions):

    os.makedirs(
        OUTPUT_DIR,
        exist_ok=True
    )

    with open(
        OUTPUT_FILE,
        "w",
        encoding="utf-8"
    ) as f:

        for q in questions:

            f.write(
                q["question"] + "\n"
            )

            f.write(
                "A." + q["options"][0] + "\n"
            )

            f.write(
                "B." + q["options"][1] + "\n"
            )

            f.write(
                "C." + q["options"][2] + "\n"
            )

            f.write(
                "D." + q["options"][3] + "\n"
            )

            f.write(
                q["answer"] + "\n\n"
            )


# ==================================================
# 主程序
# ==================================================

def main():

    print("=" * 45)

    print(
        "QuizFlow 多文件题库导入器"
    )

    print("=" * 45)

    questions = load_all_input_files()

    save_questions(questions)

    print()

    print("导入完成！")

    print(
        "共生成/导入 {} 道题。"
        .format(len(questions))
    )

    print()

    print("统一题库：")

    print(OUTPUT_FILE)


if __name__ == "__main__":
    main()